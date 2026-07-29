from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain.time_values import format_minutes

HEADER_FILL = PatternFill("solid", fgColor="2457A6")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OVERRIDE_FILL = PatternFill("solid", fgColor="FFF3D6")
SUBTLE_BORDER = Border(bottom=Side(style="thin", color="D8DEE8"))


def _build_workbook(
    title: str, headers: list[str], rows: list[list[Any]], overrides: list[set[int]]
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.sheet_view.showGridLines = False
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions

    for row_index, columns in enumerate(overrides, start=2):
        for column_index in columns:
            sheet.cell(row=row_index, column=column_index).fill = OVERRIDE_FILL

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = SUBTLE_BORDER
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"

    for column_index, header in enumerate(headers, start=1):
        longest = max(
            len(str(header)),
            *(
                len(str(sheet.cell(r, column_index).value or ""))
                for r in range(2, sheet.max_row + 1)
            ),
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(longest + 2, 12), 24
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def export_base_payroll(results: list[dict[str, Any]]) -> BytesIO:
    headers = [
        "Employee Code",
        "Employee Name",
        "Shift Hours",
        "Salary",
        "Per Day",
        "Absent",
        "Holiday",
        "Present",
        "Absent Deduction",
        "Total Hours",
        "Actual Hours",
        "Advance",
        "Late Time",
        "LT Fine",
        "Deduction",
        "Net Amount",
    ]
    rows = []
    override_columns = []
    field_columns = {
        "shiftHours": 3,
        "salary": 4,
        "perDay": 5,
        "totalAbsent": 6,
        "holiday": 7,
        "totalPresent": 8,
        "absentDeduction": 9,
        "totalHoursMinutes": 10,
        "actualHoursWorkedMinutes": 11,
        "advance": 12,
        "lateTimeMinutes": 13,
        "lateFine": 14,
        "deduction": 15,
        "netAmount": 16,
    }
    for result in results:
        rows.append(
            [
                result["employeeCode"],
                result["employeeName"],
                result["shiftHours"],
                result["salary"],
                result["perDay"],
                result["totalAbsent"],
                result["holiday"],
                result["totalPresent"],
                result["absentDeduction"],
                format_minutes(result["totalHoursMinutes"]),
                format_minutes(result["actualHoursWorkedMinutes"]),
                result["advance"],
                format_minutes(result["lateTimeMinutes"]),
                result["lateFine"],
                result["deduction"],
                result["netAmount"],
            ]
        )
        override_columns.append(
            {field_columns[field] for field in result["overriddenFields"] if field in field_columns}
        )
    return _build_workbook("Base Payroll", headers, rows, override_columns)


def export_overtime_payroll(results: list[dict[str, Any]]) -> BytesIO:
    headers = [
        "Employee Code",
        "Employee Name",
        "Shift Hours",
        "New Salary",
        "Absent",
        "Holiday",
        "Present",
        "Total Hours",
        "Per Day",
        "Actual Hours",
        "Overtime Hours",
        "OT Divisor",
        "Overtime Total",
        "Tiffin",
        "Net Amount OT",
    ]
    rows = []
    override_columns = []
    field_columns = {
        "shiftHours": 3,
        "newSalary": 4,
        "totalAbsent": 5,
        "holiday": 6,
        "totalPresent": 7,
        "totalHoursMinutes": 8,
        "perDay": 9,
        "actualHoursWorkedMinutes": 10,
        "overtimeHoursMinutes": 11,
        "overtimeShiftDivisor": 12,
        "overtimeTotal": 13,
        "tiffin": 14,
        "netAmountOt": 15,
    }
    for result in results:
        rows.append([
            result["employeeCode"],
            result["employeeName"],
            result["shiftHours"],
            result["newSalary"],
            result["totalAbsent"],
            result["holiday"],
            result["totalPresent"],
            format_minutes(result["totalHoursMinutes"]),
            result["perDay"],
            format_minutes(result["actualHoursWorkedMinutes"]),
            format_minutes(result["overtimeHoursMinutes"]),
            result["overtimeShiftDivisor"],
            result["overtimeTotal"],
            result["tiffin"],
            result["netAmountOt"],
        ])
        override_columns.append(
            {
                field_columns[field]
                for field in result.get("overriddenFields", [])
                if field in field_columns
            }
        )
    return _build_workbook("Overtime Payroll", headers, rows, override_columns)
